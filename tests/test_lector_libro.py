"""Lector de libros: elige la hoja de datos, salta títulos y excluye fórmulas sin valor (caso Hincha Store, sintético)."""
import os, sys, io
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from services import lector_libro as L

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')


def _libro_tipo_hincha():
    """Réplica estructural (datos inventados): hoja activa = Resumen de fórmulas, datos en 'Datos',
    movimientos en 'Caja' con título arriba, 'Notas' con texto. Sin valores calculados (como openpyxl)."""
    wb = openpyxl.Workbook()
    res = wb.active; res.title = 'Resumen'
    res['A1'] = 'Seguimiento financiero'; res['B4'] = 'Fecha de corte'; res['C4'] = '2026-09-29'
    res['B6'] = 'Concepto'; res['C6'] = 'Valor'; res['B7'] = 'Caja actual'; res['C7'] = '=Caja!B5'
    res['B8'] = 'Cartera'; res['C8'] = '=SUM(Datos!K2:K100)'; res['B9'] = 'Stock'; res['C9'] = '=COUNTIF(Datos!M2:M100,"Stock")'
    d = wb.create_sheet('Datos')
    d.append(['Lote', 'Equipo', 'Talla', 'Costo total (COP)', 'Precio de lista (COP)', 'Cobrado (COP)', 'Por cobrar (COP)',
              'Cliente', 'Estado (original)', 'Estado analítico', 'Utilidad a lista (COP)'])
    estados = ['Vendidada', 'Debe', 'Stock', 'Perdida']
    for i in range(1, 21):
        est = estados[i % 4]
        d.append(['L-001', f'Equipo {i % 5}', 'L', 50000, 80000, 70000 if est == 'Vendidada' else 0,
                  80000 if est == 'Debe' else 0, f'Cliente {i}' if est != 'Stock' else None, est,
                  f'=IF(I{i + 1}="Vendidada","Vendida","Otro")', f'=E{i + 1}-D{i + 1}'])
    c = wb.create_sheet('Caja')
    c['A1'] = 'Caja y movimientos'; c['A2'] = 'Registra aquí cada cobro'
    c['A4'] = 'Caja al corte'; c['B4'] = 740000
    c.append([]); c.append([])
    c['A8'], c['B8'], c['C8'], c['D8'], c['E8'] = 'Fecha', 'Tipo', 'Monto (COP)', 'Nota', 'Efecto en caja (COP)'
    c['A9'], c['B9'], c['C9'], c['D9'], c['E9'] = '2026-10-06', 'Cobro de cliente', 85000, 'abono', '=C9'
    n = wb.create_sheet('Notas')
    for i, t in enumerate(['Notas y supuestos', 'Vendida = pagada.', 'Fiada = con saldo.'], 1):
        n.cell(i, 1, t)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_clasifica_hojas_y_elige_datos():
    b = _libro_tipo_hincha()
    a = L.analizar_libro(b)
    roles = {h['nombre']: h['rol'] for h in a['hojas']}
    assert roles == {'Resumen': 'DERIVADA', 'Datos': 'DATOS', 'Caja': 'MOVIMIENTOS', 'Notas': 'NOTAS'}, roles
    assert a['hoja_sugerida'] == 'Datos'          # la hoja activa era 'Resumen'


def test_excluye_columnas_formula_sin_valor():
    m, info = L.matriz_de_hoja(_libro_tipo_hincha())
    assert info['hoja_usada'] == 'Datos' and info['filas'] == 20
    assert 'Estado analítico' not in m[0] and 'Utilidad a lista (COP)' not in m[0]
    assert {e['columna'] for e in info['columnas_excluidas']} == {'Estado analítico', 'Utilidad a lista (COP)'}
    assert m[0][:3] == ['Lote', 'Equipo', 'Talla'] and len(m) == 21
    assert all(len(f) == len(m[0]) for f in m)


def test_usuario_puede_forzar_otra_hoja():
    m, info = L.matriz_de_hoja(_libro_tipo_hincha(), 'Caja')
    assert info['hoja_usada'] == 'Caja' and info['fila_encabezado'] == 8
    assert m[0][:3] == ['Fecha', 'Tipo', 'Monto (COP)']


def test_salta_fila_de_titulo_y_celdas_combinadas():
    m, info = L.matriz_de_hoja(open(os.path.join(FIX, '03_datos_sucios_edge_cases.xlsx'), 'rb').read())
    assert info['fila_encabezado'] == 3 and m[0][0] == 'Fecha Doc' and info['filas'] == 8   # filas vacías omitidas
    assert m[1][1] == 'Camiseta Deportiva Azul'                                              # espacios recortados


def test_libro_de_produccion_trae_dos_hojas_de_datos():
    a = L.analizar_libro(open(os.path.join(FIX, '02_modelo_produccion_recetas.xlsx'), 'rb').read())
    assert [h['rol'] for h in a['hojas']] == ['DATOS', 'DATOS']


def test_numeros_como_texto_limpio():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'V'
    ws.append(['Producto', 'Precio', 'Cantidad']); ws.append(['A', 18900.0, 2]); ws.append(['B', 61612.5, 1])
    buf = io.BytesIO(); wb.save(buf)
    m, _ = L.matriz_de_hoja(buf.getvalue())
    assert m[1] == ['A', '18900', '2'] and m[2] == ['B', '61612.5', '1']    # sin '18900.0' (evita ×10)


if __name__ == '__main__':
    for n, f in list(globals().items()):
        if n.startswith('test_'):
            f(); print('OK', n)
    print('LECTOR DE LIBROS: TODAS LAS PRUEBAS PASARON')
