"""Lectura inteligente de libros Excel para el importador.

Problemas que resuelve (caso real Hincha Store):
  * El libro tiene varias hojas y la activa es un resumen de fórmulas: hay que ELEGIR la hoja de datos.
  * Filas de título / celdas combinadas encima de los encabezados.
  * Columnas de fórmula sin valor guardado (libro nunca calculado en Excel): no se pueden importar.
Devuelve siempre información explicable para mostrarle al usuario qué se leyó y qué se ignoró.
"""
import io
import math
import re
import unicodedata
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

VENTANA_ENCABEZADO = 25
_NOTAS = {'notas', 'nota', 'leeme', 'readme', 'instrucciones', 'supuestos', 'ayuda', 'como usar', 'leyenda'}
_DERIVADAS = {'resumen', 'dashboard', 'tablero', 'proyeccion', 'proyecciones', 'reporte', 'informe', 'kpi', 'kpis'}
_EXCLUYE_MOV = {'cliente', 'producto', 'equipo', 'cantidad', 'sku', 'lote', 'talla', 'servicio', 'item'}


def _norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', str(s)) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip().lower()


def _es_formula(v):
    return isinstance(v, str) and v.startswith('=')


def _a_texto(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'VERDADERO' if v else 'FALSO'
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d') if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if math.isnan(v):
            return ''
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def _fila_vacia(fila):
    return not any(c is not None and str(c).strip() != '' for c in fila)


def _detectar_encabezado(filas):
    """Índice (0-based) de la fila de encabezados entre las primeras filas, saltando títulos."""
    ventana = filas[:VENTANA_ENCABEZADO]
    conteo = [sum(1 for c in f if c is not None and str(c).strip() != '') for f in ventana]
    if not conteo or max(conteo) == 0:
        return None
    umbral = max(2, math.ceil(0.6 * max(conteo)))
    for i, f in enumerate(ventana):
        if conteo[i] < umbral:
            continue
        textos = sum(1 for c in f if isinstance(c, str) and c.strip() and not _es_formula(c)
                     and not re.fullmatch(r'[\d.,\s$%-]+', c.strip()))
        if textos / conteo[i] >= 0.7:
            return i
    return next((i for i, n in enumerate(conteo) if n > 0), None)


def _encabezados_unicos(fila, filas_datos):
    nombres, usados = [], {}
    for j, c in enumerate(fila):
        nombre = _a_texto(c)
        if not nombre:
            if any(j < len(r) and r[j] is not None and str(r[j]).strip() for r in filas_datos[:200]):
                nombre = f"Columna {get_column_letter(j + 1)}"
            else:
                nombres.append(None)
                continue
        n = usados.get(nombre, 0)
        usados[nombre] = n + 1
        nombres.append(nombre if n == 0 else f"{nombre} ({n + 1})")
    return nombres


def _analizar_hoja(ws_f, ws_v):
    filas_f = [list(r) for r in ws_f.iter_rows(values_only=True)]
    filas_v = [list(r) for r in ws_v.iter_rows(values_only=True)]
    ancho = max((len(r) for r in filas_f), default=0)
    for lst in (filas_f, filas_v):
        for r in lst:
            r.extend([None] * (ancho - len(r)))
    ih = _detectar_encabezado(filas_f)
    base = {'nombre': ws_f.title, 'filas': 0, 'columnas': 0, 'fila_encabezado': None, 'rol': 'VACIA',
            'pct_formulas': 0.0, 'columnas_formula': [], 'columnas_formula_sin_valor': [], 'encabezados': []}
    if ih is None:
        return base, None
    cuerpo_f = [(i, r) for i, r in enumerate(filas_f[ih + 1:], ih + 1) if not _fila_vacia(filas_v[i]) or not _fila_vacia(r)]
    nombres = _encabezados_unicos(filas_f[ih], [r for _, r in cuerpo_f])
    cols = [j for j, n in enumerate(nombres) if n]
    n_ne = n_form = 0
    col_formula, col_sin_valor = [], []
    for j in cols:
        ne = sum(1 for _, r in cuerpo_f if r[j] is not None and str(r[j]).strip() != '')
        fo = sum(1 for _, r in cuerpo_f if _es_formula(r[j]))
        con_valor = sum(1 for i, r in cuerpo_f if _es_formula(r[j]) and filas_v[i][j] is not None)
        n_ne += ne
        n_form += fo
        if ne and fo >= 0.5 * ne:
            col_formula.append(nombres[j])
            if con_valor < 0.5 * fo:
                col_sin_valor.append(nombres[j])
    pct = round(n_form / n_ne, 3) if n_ne else 0.0
    toks = set(_norm(' '.join(nombres[j] for j in cols)).replace('(', ' ').replace(')', ' ').split())
    nombre_hoja = _norm(ws_f.title)
    if not cuerpo_f:
        rol = 'VACIA'
    elif nombre_hoja in _NOTAS or len(cols) <= 1:
        rol = 'NOTAS'
    elif ({'fecha'} & toks and {'monto', 'valor', 'importe'} & toks and {'tipo', 'concepto', 'nota', 'descripcion'} & toks
          and not (_EXCLUYE_MOV & toks)):
        rol = 'MOVIMIENTOS'
    elif pct >= 0.35 or (nombre_hoja in _DERIVADAS and pct >= 0.2):
        rol = 'DERIVADA'
    else:
        rol = 'DATOS'
    info = dict(base, filas=len(cuerpo_f), columnas=len(cols), fila_encabezado=ih + 1, rol=rol, pct_formulas=pct,
                columnas_formula=col_formula, columnas_formula_sin_valor=col_sin_valor,
                encabezados=[nombres[j] for j in cols])
    return info, (ih, nombres, cols, cuerpo_f, filas_v)


def _abrir(file_bytes):
    wb_f = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    wb_v = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    return wb_f, wb_v


def _sugerida(hojas):
    datos = [h for h in hojas if h['rol'] == 'DATOS']
    if datos:
        return max(datos, key=lambda h: h['filas'] * max(h['columnas'], 1))['nombre']
    otras = [h for h in hojas if h['rol'] not in ('VACIA', 'NOTAS')]
    return otras[0]['nombre'] if otras else None


def analizar_libro(file_bytes):
    """Lista de hojas con su rol (DATOS / DERIVADA / MOVIMIENTOS / NOTAS / VACIA) y la hoja sugerida."""
    wb_f, wb_v = _abrir(file_bytes)
    hojas = [_analizar_hoja(wb_f[n], wb_v[n])[0] for n in wb_f.sheetnames]
    return {'hojas': hojas, 'hoja_sugerida': _sugerida(hojas)}


def matriz_de_hoja(file_bytes, hoja=None):
    """(filas_matriz, info). filas_matriz[0] = encabezados; el resto, texto limpio.
    Excluye columnas de fórmula sin valor guardado y lo informa en info['columnas_excluidas']."""
    wb_f, wb_v = _abrir(file_bytes)
    analisis = [_analizar_hoja(wb_f[n], wb_v[n]) for n in wb_f.sheetnames]
    hojas = [a[0] for a in analisis]
    sugerida = _sugerida(hojas)
    elegida = hoja if hoja in wb_f.sheetnames else sugerida
    info = {'hojas': hojas, 'hoja_sugerida': sugerida, 'hoja_usada': elegida, 'fila_encabezado': None,
            'columnas_formula': [], 'columnas_excluidas': [], 'filas': 0}
    if not elegida:
        return [], info
    h, detalle = next(a for a in analisis if a[0]['nombre'] == elegida)
    if detalle is None:
        return [], info
    ih, nombres, cols, cuerpo_f, filas_v = detalle
    excluidas = set(h['columnas_formula_sin_valor'])
    cols_out = [j for j in cols if nombres[j] not in excluidas]
    matriz = [[nombres[j] for j in cols_out]]
    for i, _ in cuerpo_f:
        matriz.append([_a_texto(filas_v[i][j]) for j in cols_out])
    info.update(fila_encabezado=ih + 1, filas=len(cuerpo_f),
                columnas_formula=[c for c in h['columnas_formula'] if c not in excluidas],
                columnas_excluidas=[{'columna': c, 'motivo': 'FORMULA_SIN_VALOR'} for c in h['columnas_formula_sin_valor']])
    return matriz, info
